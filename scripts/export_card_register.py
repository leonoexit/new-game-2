"""Export the card register workbook as read-only data for the local devtool.

Requires openpyxl. Run from any directory:
    python3 scripts/export_card_register.py
    python3 scripts/export_card_register.py --check
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/current/LITTLE-VALLEY-CARD-REGISTER-V0.xlsx"
OUTPUT = ROOT / "prototype/card-register-data.json"
HEADER_ROW = 6

# Exact assets used by prototype/app.mjs; no filename guessing or implied art approval.
RUNTIME_ART = {
    "Soil (Empty Soil)": [("Empty Soil", "empty-soil-512-v0.1.png")],
    "Soil (Overgrown Soil)": [("Overgrown Soil", "forgotten-garden-512-v0.1.png")],
    "Soil (Tilled Soil)": [("Tilled Soil", "turned-garden-512-v0.1.png")],
    "Field Rock": [("Blocked plot", "field-rock-512-v0.1.png")],
    "Turnip Crop (Growing)": [("Growing", "turnip-growing-512-v0.1.png")],
    "Turnip Crop (Watered)": [("Watered", "turnip-watered-512-v0.1.png")],
    "Turnip Crop (Mature)": [("Mature", "turnip-mature-512-v0.1.png")],
    "Runner Bean Crop (Growing)": [("Growing", "young-bean-trellis-512-v0.2.png")],
    "Runner Bean Crop (Watered)": [("Watered · shared runtime image", "young-bean-trellis-512-v0.2.png")],
}


def cell_text(value: object) -> str:
    return "" if value is None else str(value)


def build_data() -> dict:
    sheet = load_workbook(SOURCE, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [cell_text(value) for value in rows[HEADER_ROW - 1]]
    if headers[:5] != ["Identity", "Type", "State", "Action", "Description"] or headers[-1] != "Base identity" or len(headers) != 15:
        raise ValueError("Unexpected card register header; review exporter before updating")
    cards = []
    blank_rows = []
    for number, row in enumerate(rows[HEADER_ROW:], start=HEADER_ROW + 1):
        values = [cell_text(value) for value in row[:len(headers)]]
        values += [""] * (len(headers) - len(values))
        if not any(values):
            blank_rows.append(number)
            continue
        if not values[0]:
            raise ValueError(f"Register row {number} has content but no identity")
        if not values[14]:
            raise ValueError(f"Register row {number} has no base identity")
        if values[0] != values[14] and (not values[2] or values[0] != f"{values[14]} ({values[2]})"):
            raise ValueError(f"State row {number} has mismatched identity/state label")
        art = [{"label": label, "path": f"./assets/{filename}"}
               for label, filename in RUNTIME_ART.get(values[0], [])]
        for variant in art:
            if not (ROOT / "prototype" / variant["path"]).is_file():
                raise FileNotFoundError(f"Runtime art missing: {variant['path']}")
        entry_id = "entry-" + hashlib.sha1(values[0].encode("utf-8")).hexdigest()[:12]
        cards.append({"id": entry_id, "row": number, "values": values, "art": art})
    if len(cards) != 113 or blank_rows != [14]:
        raise ValueError(f"Unexpected register shape: {len(cards)} cards; blank rows {blank_rows}")
    if len({card["id"] for card in cards}) != len(cards):
        raise ValueError("Duplicate register identity/state labels")
    return {
        "source": "docs/current/LITTLE-VALLEY-CARD-REGISTER-V0.xlsx",
        "sheet": sheet.title,
        "sourceSha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "headers": headers,
        "blankRows": blank_rows,
        "cards": cards,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="fail if exported JSON is stale")
    args = parser.parse_args()
    rendered = json.dumps(build_data(), ensure_ascii=False, indent=2) + "\n"
    if args.check:
        if not OUTPUT.exists() or OUTPUT.read_text(encoding="utf-8") != rendered:
            raise SystemExit("Card register export is stale; run python3 scripts/export_card_register.py")
        print("Card register export is current: 113 entries, 15 fields")
    else:
        OUTPUT.write_text(rendered, encoding="utf-8")
        print(f"Wrote {OUTPUT.relative_to(ROOT)}: 113 entries, 15 fields")


if __name__ == "__main__":
    main()
